import ast
import unittest
from pathlib import Path

import openpyxl
from flask import Flask, render_template_string

from core.auditoria import build_reporte_gerencial, marcar_resultado_revisado
from core.auditoria_export import generar_excel_auditoria
from core.models import (
    AjusteInventario,
    AuditoriaResultado,
    ConteoDetalle,
    InventarioPeriodo,
    Justificacion,
    Producto,
    db,
)
from core.periodos import registrar_conteo_admin
from core.security import rol_permitido


CLIENTE = "SEG"
TIENDA = "T-SEG"


class PruebasPermisosSeguridad(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = Flask(__name__)
        cls.app.config.update(
            TESTING=True,
            SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(cls.app)
        cls.contexto = cls.app.app_context()
        cls.contexto.push()

    @classmethod
    def tearDownClass(cls):
        db.session.remove()
        cls.contexto.pop()

    def setUp(self):
        db.session.remove()
        db.drop_all()
        db.create_all()

    def tearDown(self):
        db.session.rollback()
        db.session.remove()

    def crear_periodo(self, estado="Cerrado"):
        periodo = InventarioPeriodo(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            numero=1,
            fecha_desde="2026-08-01",
            fecha_hasta="2026-08-08",
            dias_periodo=7,
            estado=estado,
            usuario_creador="admin",
        )
        db.session.add(periodo)
        db.session.flush()
        return periodo

    def test_empleado_no_tiene_permiso_y_rutas_sensibles_exigen_admin(self):
        self.assertFalse(rol_permitido("empleado", "administrador"))
        self.assertTrue(rol_permitido("administrador", "administrador"))

        codigo = Path(__file__).resolve().parents[1].joinpath("app.py").read_text(encoding="utf-8")
        arbol = ast.parse(codigo)
        rutas_sensibles = {
            "admin_periodo_excel_importar",
            "admin_excel_vincular",
            "admin_auditoria_ejecutar",
            "admin_auditoria",
            "admin_asistente_consultar",
            "admin_justificar",
            "admin_marcar_revisado",
            "admin_reporte_gerencial",
            "admin_auditoria_exportar",
        }
        funciones = {
            nodo.name: nodo
            for nodo in ast.walk(arbol)
            if isinstance(nodo, (ast.FunctionDef, ast.AsyncFunctionDef))
        }
        for nombre in rutas_sensibles:
            decoradores = funciones[nombre].decorator_list
            protegido = any(
                isinstance(decorador, ast.Call)
                and isinstance(decorador.func, ast.Name)
                and decorador.func.id == "login_required"
                and any(
                    kw.arg == "rol"
                    and isinstance(kw.value, ast.Constant)
                    and kw.value.value == "administrador"
                    for kw in decorador.keywords
                )
                for decorador in decoradores
            )
            self.assertTrue(protegido, f"{nombre} debe exigir rol administrador")

    def test_admin_corrige_periodo_historico_mediante_ajuste_trazable(self):
        periodo = self.crear_periodo("Cerrado")
        conteo = ConteoDetalle(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            usuario="empleado",
            producto_nombre="Producto histórico",
            categoria="Seguridad",
            cantidad_unidad=5,
            total_unidad_base=5,
            fue_cargado=True,
        )
        db.session.add(conteo)
        db.session.flush()

        tipo, ajuste = registrar_conteo_admin(
            periodo,
            producto_nombre=conteo.producto_nombre,
            categoria=conteo.categoria,
            cantidad=8,
            usuario="admin-seguro",
        )
        db.session.flush()

        self.assertEqual(tipo, "ajuste")
        self.assertEqual(conteo.total_unidad_base, 5)
        self.assertEqual(ajuste.cantidad_ajustada, 3)
        self.assertEqual(ajuste.usuario_admin, "admin-seguro")
        self.assertEqual(ajuste.motivo, "Corrección de carga")
        self.assertEqual(AjusteInventario.query.count(), 1)

    def test_marcar_revisado_crea_trazabilidad_del_admin(self):
        periodo = self.crear_periodo("Conciliado")
        resultado = AuditoriaResultado(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            producto_nombre="Producto revisado",
            diferencia=-2,
            tipo_diferencia="faltante",
            estado_auditoria="Pendiente",
        )
        db.session.add(resultado)
        db.session.flush()

        traza = marcar_resultado_revisado(resultado, "admin-auditor")
        db.session.flush()

        self.assertEqual(resultado.estado_auditoria, "Revisado")
        self.assertEqual(traza.usuario, "admin-auditor")
        self.assertEqual(traza.causa, "Marcado como revisado")
        self.assertEqual(Justificacion.query.filter_by(resultado_id=resultado.id).count(), 1)

    def test_datos_maliciosos_no_inyectan_sql_html_ni_formulas_excel(self):
        nombre = '=HYPERLINK("https://malicioso") Robert\'); DROP TABLE productos;-- <script>alert(1)</script> ñ'
        periodo = self.crear_periodo("Auditado")
        producto = Producto(nombre=nombre, categoria="Raro & especial", visible_empleado=True)
        db.session.add(producto)
        db.session.flush()
        resultado = AuditoriaResultado(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            producto_nombre=nombre,
            categoria=producto.categoria,
            articulo_codigo="+CMD|'/C calc'!A0",
            diferencia=-1,
            tipo_diferencia="faltante",
            severidad="Observación",
            impacto=100,
            causa_sugerida="Pendiente de revisión & <b>urgente</b>",
            evidencia="Texto con comillas ' \" ; -- y Unicode: áéíóú",
            estado_auditoria="Pendiente",
        )
        db.session.add(resultado)
        db.session.flush()
        db.session.add(Justificacion(
            resultado_id=resultado.id,
            cliente_id=CLIENTE,
            causa="Prueba especial",
            cantidad_justificada=1,
            observacion="@SUM(1,1) <script>otro</script>",
            usuario="admin",
        ))
        db.session.commit()

        self.assertEqual(Producto.query.filter_by(nombre=nombre).count(), 1)
        self.assertGreaterEqual(Producto.query.count(), 1)
        reporte = build_reporte_gerencial(periodo)
        self.assertEqual(reporte["faltantes"][0].producto_nombre, nombre)

        html = render_template_string("{{ nombre }}", nombre=nombre)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)

        libro = openpyxl.load_workbook(generar_excel_auditoria(periodo), data_only=False)
        hoja = libro.active
        self.assertEqual(hoja.cell(2, 5).data_type, "s")
        self.assertTrue(hoja.cell(2, 5).value.startswith("'="))
        self.assertEqual(hoja.cell(2, 4).data_type, "s")
        self.assertTrue(hoja.cell(2, 4).value.startswith("'+"))
        encabezados = [celda.value for celda in hoja[1]]
        columna_observacion = encabezados.index("Observación") + 1
        self.assertEqual(hoja.cell(2, columna_observacion).data_type, "s")
        self.assertTrue(hoja.cell(2, columna_observacion).value.startswith("'@"))
        libro.close()


if __name__ == "__main__":
    unittest.main()
