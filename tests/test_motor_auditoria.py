import unittest

import openpyxl
from flask import Flask

from core.auditoria import ejecutar_auditoria
from core.auditoria_export import generar_excel_auditoria
from core.ajustes import ajuste_es_baja_no_imputable
from core.models import (
    AjusteInventario,
    AsistenteIAConsulta,
    AuditoriaResultado,
    ConteoDetalle,
    ExcelDetalle,
    ExcelImportado,
    InventarioPeriodo,
    Justificacion,
    Producto,
    ProductoPrecio,
    RegistroAveriado,
    RegistroVencimiento,
    db,
)


CLIENTE = "TEST"
TIENDA = "TTEST"


class PruebasMotorAuditoria(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = Flask(__name__)
        cls.app.config.update(
            TESTING=True,
            SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(cls.app)
        cls.contexto = cls.app.app_context()
        cls.contexto.push()

    @classmethod
    def tearDownClass(cls):
        db.session.remove()
        cls.contexto.pop()

    def setUp(self):
        db.session.remove()
        db.drop_all()
        db.create_all()

    def tearDown(self):
        db.session.rollback()
        db.session.remove()

    def crear_producto(self, nombre):
        producto = Producto(
            nombre=nombre,
            categoria="Pruebas",
            visible_empleado=True,
        )
        db.session.add(producto)
        db.session.flush()
        return producto

    def crear_periodo(
        self,
        numero,
        desde="2026-08-01",
        hasta="2026-08-08",
        estado="Cerrado",
    ):
        periodo = InventarioPeriodo(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            numero=numero,
            fecha_desde=desde,
            fecha_hasta=hasta,
            dias_periodo=7,
            estado=estado,
            usuario_creador="tester",
        )
        db.session.add(periodo)
        db.session.flush()
        return periodo

    def agregar_conteo(self, periodo, producto, cantidad):
        conteo = ConteoDetalle(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            usuario="empleado",
            producto_nombre=producto.nombre,
            categoria=producto.categoria,
            cantidad_unidad=cantidad,
            cantidad_caja=0,
            cantidad_bulto=0,
            total_unidad_base=cantidad,
            fue_cargado=True,
        )
        db.session.add(conteo)
        db.session.flush()
        return conteo

    def agregar_excel(
        self,
        periodo,
        producto,
        *,
        stock_inicial=0,
        compras=0,
        otros_ingresos=0,
        ventas=0,
        otras_salidas=0,
        stock_final=0,
        costo=None,
        excluido=False,
    ):
        excel = ExcelImportado.query.filter_by(periodo_id=periodo.id).first()
        if excel is None:
            excel = ExcelImportado(
                periodo_id=periodo.id,
                cliente_id=CLIENTE,
                nombre_archivo=f"periodo-{periodo.numero}.xlsx",
                usuario_importador="tester",
                estado_validacion="ok",
            )
            db.session.add(excel)
            db.session.flush()
        detalle = ExcelDetalle(
            excel_id=excel.id,
            articulo=f"P-{producto.id}",
            artdescrip=producto.nombre,
            stockinicial=stock_inicial,
            compras=compras,
            otrosingresos=otros_ingresos,
            ventareal=ventas,
            otrassalidas=otras_salidas,
            stockfinal=stock_final,
            producto_nombre_interno=None if excluido else producto.nombre,
            producto_id=None if excluido else producto.id,
            estado_vinculacion="excluido" if excluido else "vinculado",
            excluido_auditoria=excluido,
            motivo_exclusion="Excluido en prueba" if excluido else None,
        )
        db.session.add(detalle)
        if costo is not None:
            precio = ProductoPrecio.query.filter_by(producto_id=producto.id).first()
            if precio is None:
                precio = ProductoPrecio(
                    cliente_id=CLIENTE,
                    producto_id=producto.id,
                    producto_nombre=producto.nombre,
                    categoria=producto.categoria,
                )
                db.session.add(precio)
            precio.precio = costo
        db.session.flush()
        return detalle

    def test_producto_excluido_no_se_calcula_ni_se_exporta(self):
        periodo = self.crear_periodo(1)
        producto = self.crear_producto("Producto excluido")
        producto.codigo_articulo = f"P-{producto.id}"
        self.agregar_conteo(periodo, producto, 9)
        detalle = self.agregar_excel(
            periodo, producto, stock_inicial=20, compras=5, ventas=3,
            stock_final=9, excluido=True,
        )
        detalle.artdescrip = "Descripción externa excluida"
        db.session.commit()

        resultados = ejecutar_auditoria(periodo)
        db.session.commit()
        self.assertEqual(resultados, [])
        self.assertEqual(AuditoriaResultado.query.filter_by(periodo_id=periodo.id).count(), 0)

        # Simula un resultado antiguo creado antes de la corrección. La defensa
        # del exportador también debe omitirlo.
        db.session.add(AuditoriaResultado(
            periodo_id=periodo.id, cliente_id=CLIENTE,
            producto_nombre=producto.nombre, articulo_codigo=f"P-{producto.id}",
        ))
        db.session.commit()
        libro = openpyxl.load_workbook(generar_excel_auditoria(periodo), data_only=True)
        self.assertEqual(libro.active.max_row, 1)
        libro.close()

    def test_reejecutar_conserva_resultado_y_trazabilidad(self):
        periodo = self.crear_periodo(1)
        producto = self.crear_producto("Producto con trazabilidad")
        self.agregar_conteo(periodo, producto, 5)
        detalle = self.agregar_excel(
            periodo, producto, stock_inicial=10, ventas=3, stock_final=5,
        )
        resultado_inicial = ejecutar_auditoria(periodo)[0]
        db.session.commit()
        resultado_id = resultado_inicial.id
        diferencia_inicial = resultado_inicial.diferencia

        justificacion = Justificacion(
            resultado_id=resultado_id,
            cliente_id=CLIENTE,
            causa="Error de conteo",
            cantidad_justificada=1,
            importe_justificado=0,
            observacion="Trazabilidad previa a la re-ejecucion.",
            usuario="tester",
        )
        consulta = AsistenteIAConsulta(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            periodo_id=periodo.id,
            resultado_id=resultado_id,
            usuario="tester",
            pregunta="Explica el resultado",
            respuesta="Respuesta de prueba",
        )
        db.session.add_all([justificacion, consulta])
        detalle.ventareal = 4
        db.session.commit()

        resultado_actualizado = ejecutar_auditoria(periodo)[0]
        db.session.commit()

        self.assertEqual(resultado_actualizado.id, resultado_id)
        self.assertNotEqual(resultado_actualizado.diferencia, diferencia_inicial)
        self.assertEqual(justificacion.resultado_id, resultado_id)
        self.assertEqual(consulta.resultado_id, resultado_id)
        self.assertEqual(Justificacion.query.count(), 1)
        self.assertEqual(AsistenteIAConsulta.query.count(), 1)

    def test_reejecutar_archiva_producto_retirado_sin_romper_su_traza(self):
        periodo = self.crear_periodo(1)
        producto = self.crear_producto("Producto retirado")
        self.agregar_conteo(periodo, producto, 2)
        detalle = self.agregar_excel(
            periodo, producto, stock_inicial=8, ventas=4, stock_final=2,
        )
        resultado = ejecutar_auditoria(periodo)[0]
        db.session.commit()
        db.session.add(Justificacion(
            resultado_id=resultado.id,
            cliente_id=CLIENTE,
            causa="Pendiente de revision",
            cantidad_justificada=0,
            importe_justificado=0,
            observacion="Conservar aunque el producto sea retirado.",
            usuario="tester",
        ))
        detalle.excluido_auditoria = True
        detalle.estado_vinculacion = "excluido"
        db.session.commit()

        self.assertEqual(ejecutar_auditoria(periodo), [])
        db.session.commit()
        archivado = db.session.get(AuditoriaResultado, resultado.id)
        self.assertEqual(archivado.estado_auditoria, "Archivado")
        self.assertEqual(archivado.impacto, 0)
        self.assertEqual(Justificacion.query.count(), 1)
        libro = openpyxl.load_workbook(generar_excel_auditoria(periodo), data_only=True)
        self.assertEqual(libro.active.max_row, 1)
        libro.close()

    def agregar_resultado_anterior(
        self,
        periodo,
        producto,
        *,
        conteo_final=10,
        diferencia=0,
        compras=0,
    ):
        resultado = AuditoriaResultado(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            producto_nombre=producto.nombre,
            categoria=producto.categoria,
            conteo_final=conteo_final,
            diferencia=diferencia,
            compras=compras,
        )
        db.session.add(resultado)
        db.session.flush()
        return resultado

    def resultado_unico(self, periodo):
        resultados = ejecutar_auditoria(periodo)
        self.assertEqual(len(resultados), 1)
        return resultados[0]

    def test_formula_stock_esperado(self):
        producto = self.crear_producto("Stock esperado")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 86)
        self.agregar_excel(
            periodo,
            producto,
            stock_inicial=100,
            compras=20,
            otros_ingresos=5,
            ventas=30,
            otras_salidas=4,
            stock_final=86,
        )
        db.session.add(RegistroAveriado(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            fecha="2026-08-03",
            usuario="empleado",
            categoria="Pruebas",
            producto=producto.nombre,
            cantidad=3,
            cantidad_unidades=3,
            sinc_estado="sincronizado",
        ))
        db.session.add(RegistroVencimiento(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            fecha="2026-08-04",
            usuario="empleado",
            categoria="Pruebas",
            producto=producto.nombre,
            cantidad=2,
            cantidad_unidades=2,
            fecha_vencimiento="2026-08-10",
            sinc_estado="sincronizado",
        ))
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.stock_esperado, 86)
        self.assertEqual(resultado.venta_teorica, 30)
        self.assertEqual(resultado.diferencia, 0)

    def test_formula_oficial_venta_teorica_menos_venta_real(self):
        producto = self.crear_producto("Bombon suizo x unidad")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 5)
        self.agregar_excel(
            periodo,
            producto,
            stock_inicial=136,
            compras=48,
            ventas=256,
            costo=6500,
        )

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.venta_teorica, 179)
        self.assertEqual(resultado.diferencia, -77)
        self.assertEqual(resultado.tipo_diferencia, "sobrante")
        self.assertEqual(resultado.impacto, 500500)

    def test_formula_conteo_final(self):
        producto = self.crear_producto("Conteo final")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 80)
        self.agregar_excel(periodo, producto, stock_inicial=85)
        db.session.add(AjusteInventario(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            producto_nombre=producto.nombre,
            usuario_admin="admin",
            cantidad_ajustada=5,
            motivo="Corrección de carga",
            impacta_stock=True,
        ))
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.conteo_empleado, 80)
        self.assertEqual(resultado.ajuste_admin, 5)
        self.assertEqual(resultado.conteo_final, 85)

    def test_formula_diferencia(self):
        producto = self.crear_producto("Diferencia")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 7)
        self.agregar_excel(periodo, producto, stock_inicial=10)

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.stock_esperado, 10)
        self.assertEqual(resultado.conteo_final, 7)
        self.assertEqual(resultado.venta_teorica, 3)
        self.assertEqual(resultado.diferencia, 3)
        self.assertEqual(resultado.tipo_diferencia, "faltante")

    def test_tipos_de_diferencia_y_regla_stock_cero(self):
        periodo = self.crear_periodo(1)
        casos = [
            ("Positiva", 10, 12, "sobrante"),
            ("Negativa", 10, 8, "faltante"),
            ("Cero", 10, 10, "correcto"),
            ("Stock esperado cero", 0, 3, "sobrante"),
        ]
        for nombre, esperado, conteo, _ in casos:
            producto = self.crear_producto(nombre)
            self.agregar_conteo(periodo, producto, conteo)
            self.agregar_excel(periodo, producto, stock_inicial=esperado)

        resultados = {
            resultado.producto_nombre: resultado
            for resultado in ejecutar_auditoria(periodo)
        }

        for nombre, _, _, tipo in casos:
            with self.subTest(nombre=nombre):
                self.assertEqual(resultados[nombre].tipo_diferencia, tipo)

    def test_limites_de_severidad(self):
        periodo = self.crear_periodo(1)
        casos = [
            ("Severidad 0", 30, "Correcto"),
            ("Severidad 5", 35, "Observación"),
            ("Severidad 20", 50, "Revisar"),
            ("Severidad 21", 51, "Crítico"),
        ]
        for nombre, conteo, _ in casos:
            producto = self.crear_producto(nombre)
            self.agregar_conteo(periodo, producto, conteo)
            self.agregar_excel(periodo, producto, stock_inicial=30)

        resultados = {
            resultado.producto_nombre: resultado
            for resultado in ejecutar_auditoria(periodo)
        }

        for nombre, _, severidad in casos:
            with self.subTest(nombre=nombre):
                self.assertEqual(resultados[nombre].severidad, severidad)

    def test_impacto_economico(self):
        producto = self.crear_producto("Impacto")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 7)
        self.agregar_excel(
            periodo,
            producto,
            stock_inicial=10,
            costo=500,
        )

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.diferencia, 3)
        self.assertEqual(resultado.costo_unitario, 500)
        self.assertEqual(resultado.impacto, 1500)
        self.assertEqual(resultado.fuente_costo, "Precio interno")

    def test_causa_compra_mal_cargada(self):
        producto = self.crear_producto("Compra sospechosa")
        anterior = self.crear_periodo(1)
        self.agregar_resultado_anterior(
            anterior, producto, conteo_final=10, compras=10
        )
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        self.agregar_conteo(actual, producto, 49)
        self.agregar_excel(actual, producto, stock_inicial=10, compras=40)

        resultado = self.resultado_unico(actual)

        self.assertEqual(resultado.causa_sugerida, "Compra mal cargada")
        self.assertEqual(resultado.nivel_confianza, "Alto")

    def test_causa_error_de_conteo(self):
        producto = self.crear_producto("Error conteo")
        anterior = self.crear_periodo(1)
        self.agregar_resultado_anterior(
            anterior, producto, conteo_final=10, diferencia=-10
        )
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        self.agregar_conteo(actual, producto, 3)
        self.agregar_excel(actual, producto, stock_inicial=10)

        resultado = self.resultado_unico(actual)

        self.assertEqual(resultado.causa_sugerida, "Error de conteo")
        self.assertEqual(resultado.nivel_confianza, "Medio")
        self.assertEqual(resultado.diferencia_anterior_compensada, -10)

    def test_merma_reduce_esperado_una_vez_y_no_genera_impacto(self):
        producto = self.crear_producto("Producto con merma")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 7)
        self.agregar_excel(periodo, producto, stock_inicial=10, costo=1000)
        db.session.add(RegistroAveriado(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            fecha="2026-08-03",
            usuario="empleado",
            categoria="Pruebas",
            producto=producto.nombre,
            cantidad=3,
            cantidad_unidades=3,
            sinc_estado="sincronizado",
        ))
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.stock_esperado, 7)
        self.assertEqual(resultado.conteo_final, 7)
        self.assertEqual(resultado.diferencia, 0)
        self.assertEqual(resultado.impacto, 0)
        self.assertEqual(resultado.causa_sugerida, "Sin diferencia")
        self.assertEqual(resultado.cantidad_merma, 3)
        self.assertEqual(resultado.nivel_confianza, "Alto")
        self.assertIn("Bajas no imputables al empleado", resultado.evidencia)
        self.assertIn("descontadas una sola vez", resultado.evidencia)

    def test_vencido_reduce_esperado_una_vez_y_no_genera_impacto(self):
        producto = self.crear_producto("Producto vencido")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 7)
        self.agregar_excel(periodo, producto, stock_inicial=10, costo=1000)
        db.session.add(RegistroVencimiento(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            fecha="2026-08-03",
            usuario="empleado",
            categoria="Pruebas",
            producto=producto.nombre,
            cantidad=3,
            cantidad_unidades=3,
            fecha_vencimiento="2026-08-10",
            sinc_estado="sincronizado",
        ))
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.stock_esperado, 7)
        self.assertEqual(resultado.diferencia, 0)
        self.assertEqual(resultado.impacto, 0)
        self.assertEqual(resultado.causa_sugerida, "Sin diferencia")
        self.assertEqual(resultado.cantidad_vencida, 3)
        self.assertEqual(resultado.nivel_confianza, "Alto")
        self.assertIn("Bajas no imputables al empleado", resultado.evidencia)

    def test_merma_no_se_reutiliza_para_justificar_diferencia_residual(self):
        producto = self.crear_producto("Producto con faltante residual")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 4)
        self.agregar_excel(periodo, producto, stock_inicial=10, costo=1000)
        db.session.add(RegistroAveriado(
            cliente_id=CLIENTE, tienda_id=TIENDA, fecha="2026-08-03",
            usuario="empleado", categoria="Pruebas", producto=producto.nombre,
            cantidad=3, cantidad_unidades=3, sinc_estado="sincronizado",
        ))
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.stock_esperado, 7)
        self.assertEqual(resultado.diferencia, 3)
        self.assertEqual(resultado.impacto, 3000)
        self.assertEqual(resultado.causa_sugerida, "Pendiente de revisión")
        self.assertNotEqual(resultado.causa_sugerida, "Merma o averiado")
        self.assertIn("diferencia residual de +3.0", resultado.evidencia)
        self.assertIn("independiente de esas bajas", resultado.evidencia)

    def test_ajuste_de_baja_ya_registrada_no_duplica_descuento(self):
        producto = self.crear_producto("Producto con ajuste duplicado")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 7)
        self.agregar_excel(periodo, producto, stock_inicial=10, costo=1000)
        db.session.add_all([
            RegistroAveriado(
                cliente_id=CLIENTE, tienda_id=TIENDA, fecha="2026-08-03",
                usuario="empleado", categoria="Pruebas", producto=producto.nombre,
                cantidad=3, cantidad_unidades=3, sinc_estado="sincronizado",
            ),
            AjusteInventario(
                periodo_id=periodo.id, cliente_id=CLIENTE,
                producto_nombre=producto.nombre, usuario_admin="admin",
                cantidad_ajustada=-3, motivo="Merma o averiado ya registrado",
                impacta_stock=True,
            ),
        ])
        db.session.flush()

        resultado = self.resultado_unico(periodo)

        self.assertTrue(ajuste_es_baja_no_imputable("Merma o averiado ya registrado"))
        self.assertEqual(resultado.ajuste_admin, 0)
        self.assertEqual(resultado.conteo_final, 7)
        self.assertEqual(resultado.diferencia, 0)
        self.assertEqual(resultado.impacto, 0)

    def test_causa_inconsistencia_de_continuidad(self):
        producto = self.crear_producto("Continuidad")
        anterior = self.crear_periodo(1)
        self.agregar_resultado_anterior(anterior, producto, conteo_final=10)
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        self.agregar_conteo(actual, producto, 9)
        self.agregar_excel(actual, producto, stock_inicial=7)

        resultado = self.resultado_unico(actual)

        self.assertTrue(resultado.alerta_continuidad)
        self.assertEqual(
            resultado.causa_sugerida,
            "Inconsistencia de continuidad",
        )
        self.assertEqual(resultado.nivel_confianza, "Alto")

    def test_sin_evidencia_queda_pendiente(self):
        producto = self.crear_producto("Sin evidencia")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 9)
        self.agregar_excel(periodo, producto, stock_inicial=10)

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.causa_sugerida, "Pendiente de revisión")
        self.assertEqual(resultado.nivel_confianza, "Bajo")
        self.assertEqual(resultado.estado_auditoria, "Pendiente")

    def test_niveles_de_confianza_alto_medio_bajo(self):
        anterior = self.crear_periodo(1)
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        casos = [
            ("Confianza alta", -10, 0, "Alto"),
            ("Confianza media", -10, 3, "Medio"),
            ("Confianza baja", 0, 9, "Bajo"),
        ]
        for nombre, diferencia_anterior, conteo_actual, _ in casos:
            producto = self.crear_producto(nombre)
            self.agregar_resultado_anterior(
                anterior,
                producto,
                conteo_final=10,
                diferencia=diferencia_anterior,
            )
            self.agregar_conteo(actual, producto, conteo_actual)
            self.agregar_excel(actual, producto, stock_inicial=10)

        resultados = {
            resultado.producto_nombre: resultado
            for resultado in ejecutar_auditoria(actual)
        }

        for nombre, _, _, confianza in casos:
            with self.subTest(nombre=nombre):
                self.assertEqual(resultados[nombre].nivel_confianza, confianza)

    def test_dif_anterior_sin_periodo_anterior(self):
        producto = self.crear_producto("Sin anterior")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 9)
        self.agregar_excel(periodo, producto, stock_inicial=10)

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.diferencia_anterior_compensada, 0)

    def test_dif_anterior_con_anterior_sin_diferencia(self):
        producto = self.crear_producto("Anterior sin diferencia")
        anterior = self.crear_periodo(1)
        self.agregar_resultado_anterior(
            anterior, producto, conteo_final=10, diferencia=0
        )
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        self.agregar_conteo(actual, producto, 9)
        self.agregar_excel(actual, producto, stock_inicial=10)

        resultado = self.resultado_unico(actual)

        self.assertEqual(resultado.diferencia_anterior_compensada, 0)

    def test_dif_anterior_con_diferencia_compensada(self):
        producto = self.crear_producto("Anterior compensado")
        anterior = self.crear_periodo(1)
        self.agregar_resultado_anterior(
            anterior, producto, conteo_final=10, diferencia=-10
        )
        actual = self.crear_periodo(
            2, desde="2026-08-09", hasta="2026-08-16"
        )
        self.agregar_conteo(actual, producto, 0)
        self.agregar_excel(actual, producto, stock_inicial=10)

        resultado = self.resultado_unico(actual)

        self.assertEqual(resultado.diferencia_anterior_compensada, -10)

    def test_dif_anterior_inicializado_en_rama_sin_diferencia(self):
        producto = self.crear_producto("Inicialización segura")
        periodo = self.crear_periodo(1)
        self.agregar_conteo(periodo, producto, 10)
        self.agregar_excel(periodo, producto, stock_inicial=10)

        resultado = self.resultado_unico(periodo)

        self.assertEqual(resultado.causa_sugerida, "Sin diferencia")
        self.assertEqual(resultado.diferencia_anterior_compensada, 0)


if __name__ == "__main__":
    unittest.main()
