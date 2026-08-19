import io
import unittest

import openpyxl
from flask import Flask, session

from core import empleado as empleado_service
from core.auditoria import (
    build_reporte_gerencial,
    ejecutar_auditoria,
    justificar_resultado,
)
from core.auditoria_export import generar_excel_auditoria
from core.excel_importer import importar_excel
from core.models import (
    AuditoriaResultado,
    ConteoDetalle,
    ExcelDetalle,
    ExcelImportado,
    InventarioItem,
    InventarioPeriodo,
    Justificacion,
    Producto,
    db,
)
from core.periodos import cerrar_periodo
from core.sync_bridge import (
    propagar_conteo_a_periodo,
    retroalimentar_periodo_desde_items,
)


CLIENTE = "TEST"
TIENDA = "TTEST"
PRODUCTO = "Producto E2E"


def crear_excel_oficial():
    headers = [
        "articulo", "artdescrip", "artcosto", "stockinicial", "compras",
        "otrosingresos", "otrassalidas", "stockfinal", "ventateorica",
        "ventareal", "diferencia", "importedesvio", "kilos", "unidades",
        "grupo", "grudescrip",
    ]
    row = [
        "E2E-001", PRODUCTO, 1000, 10, 0, 0, 0, 7, 3, 0, 3, 3000,
        0, 0, "Pruebas", "Pruebas",
    ]
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(headers)
    hoja.append(row)
    stream = io.BytesIO()
    libro.save(stream)
    libro.close()
    return stream.getvalue()


class PruebaEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = Flask(__name__)
        cls.app.config.update(
            TESTING=True,
            SECRET_KEY="testing",
            SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(cls.app)
        cls.contexto = cls.app.app_context()
        cls.contexto.push()

    @classmethod
    def tearDownClass(cls):
        db.session.remove()
        cls.contexto.pop()

    def setUp(self):
        db.session.remove()
        db.drop_all()
        db.create_all()
        self.producto = Producto(
            nombre=PRODUCTO,
            categoria="Pruebas",
            visible_empleado=True,
        )
        db.session.add(self.producto)
        db.session.commit()

    def tearDown(self):
        db.session.rollback()
        db.session.remove()

    def test_flujo_completo(self):
        # 1. Crear período y sus placeholders pendientes.
        periodo = InventarioPeriodo(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            numero=1,
            fecha_desde="2026-08-01",
            fecha_hasta="2026-08-08",
            dias_periodo=7,
            estado="Abierto",
            usuario_creador="admin",
        )
        db.session.add(periodo)
        db.session.flush()
        self.assertEqual(retroalimentar_periodo_desde_items(periodo), 0)
        placeholder = ConteoDetalle.query.filter_by(
            periodo_id=periodo.id,
            producto_nombre=PRODUCTO,
        ).one()
        self.assertFalse(placeholder.fue_cargado)

        # 2. El empleado carga 7 unidades y las guarda como inventario pendiente.
        with self.app.test_request_context():
            session.update(
                cliente_id=CLIENTE,
                tienda_id=TIENDA,
                usuario="empleado",
            )
            carrito, _ = empleado_service.add_carrito_item(
                carrito=[],
                categoria="Pruebas",
                producto=PRODUCTO,
                cantidad=7,
                ume="Unidad",
                tipo_inventario="Semanal",
                fecha="2026-08-03",
                detalle="Carga E2E",
            )
            guardados = empleado_service.build_carrito_guardado(
                carrito,
                TIENDA,
                "empleado",
            )
            db.session.flush()
        self.assertEqual(guardados, 1)
        inventario = InventarioItem.query.filter_by(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            producto=PRODUCTO,
        ).one()
        self.assertEqual(inventario.cantidad, 7)
        self.assertEqual(inventario.sinc_estado, "pendiente")

        # 3. La sincronización propaga el conteo al período y confirma el envío.
        propagados = propagar_conteo_a_periodo(
            tienda_id=TIENDA,
            cliente_id=CLIENTE,
            usuario="empleado",
            periodo_id=periodo.id,
        )
        resumen_sync = empleado_service.procesar_sincronizacion(
            cliente_id=CLIENTE,
            tienda_id=TIENDA,
            usuario="empleado",
            accion="solo_enviar",
        )
        db.session.flush()
        self.assertEqual(propagados, 1)
        self.assertEqual(resumen_sync["n_inv"], 1)
        self.assertEqual(inventario.sinc_estado, "sincronizado")
        self.assertTrue(placeholder.fue_cargado)
        self.assertEqual(placeholder.total_unidad_base, 7)
        self.assertEqual(periodo.estado, "Cargado")

        # 4. El cierre pasa porque ya no quedan productos pendientes.
        cerrado, pendientes = cerrar_periodo(periodo)
        self.assertTrue(cerrado)
        self.assertEqual(pendientes, [])
        self.assertEqual(periodo.estado, "Cerrado")

        # 5. Importar el Excel oficial y vincular automáticamente el producto.
        excel, advertencias = importar_excel(
            periodo_id=periodo.id,
            cliente_id=CLIENTE,
            usuario="admin",
            filename="oficial-e2e.xlsx",
            contenido=crear_excel_oficial(),
        )
        periodo.estado = "Excel Importado"
        db.session.flush()
        self.assertIsInstance(excel, ExcelImportado)
        self.assertEqual(excel.estado_validacion, "ok")
        self.assertEqual(advertencias, [])
        detalle_excel = ExcelDetalle.query.filter_by(excel_id=excel.id).one()
        self.assertEqual(detalle_excel.producto_id, self.producto.id)
        self.assertEqual(detalle_excel.estado_vinculacion, "vinculado")

        # 6. Ejecutar auditoría: conteo 7 frente a esperado 10 = faltante 3.
        resultados = ejecutar_auditoria(periodo)
        db.session.flush()
        self.assertEqual(len(resultados), 1)
        resultado = resultados[0]
        self.assertIsInstance(resultado, AuditoriaResultado)
        self.assertEqual(resultado.stock_esperado, 10)
        self.assertEqual(resultado.conteo_final, 7)
        self.assertEqual(resultado.diferencia, -3)
        self.assertEqual(resultado.tipo_diferencia, "faltante")
        self.assertEqual(resultado.impacto, 3000)
        self.assertEqual(periodo.estado, "Conciliado")

        # 7. Justificar el faltante y completar el período auditado.
        justificacion = justificar_resultado(
            resultado,
            periodo,
            causa="Error de conteo",
            cantidad=3,
            observacion="Conteo confirmado por el administrador.",
            usuario="admin",
        )
        db.session.flush()
        self.assertIsInstance(justificacion, Justificacion)
        self.assertEqual(justificacion.importe_justificado, 3000)
        self.assertEqual(resultado.estado_auditoria, "Justificado")
        self.assertEqual(periodo.estado, "Auditado")

        # 8. Exportar y volver a abrir el XLSX generado.
        exportado = generar_excel_auditoria(periodo)
        libro = openpyxl.load_workbook(exportado, data_only=True)
        hoja = libro.active
        headers = [celda.value for celda in hoja[1]]
        fila = [celda.value for celda in hoja[2]]
        datos = dict(zip(headers, fila))
        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(datos["Producto"], PRODUCTO)
        self.assertEqual(datos["Estado Auditoría"], "Justificado")
        self.assertEqual(datos["Justificación Manual"], "Error de conteo")
        libro.close()

        # 9. Generar el reporte gerencial del mismo período.
        reporte = build_reporte_gerencial(periodo)
        self.assertEqual(reporte["total_perdida"], 3000)
        self.assertEqual(len(reporte["faltantes"]), 1)
        self.assertEqual(reporte["faltantes"][0].producto_nombre, PRODUCTO)
        self.assertIsNone(reporte["comparacion"])

        db.session.commit()


if __name__ == "__main__":
    unittest.main()
