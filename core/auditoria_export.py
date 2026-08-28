"""Generación del archivo Excel de resultados de auditoría."""
from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    AuditoriaResultado, ExcelDetalle, ExcelImportado, InventarioPeriodo,
    Justificacion,
)


HEADERS_AUDITORIA = [
    "Inventario", "Fecha Desde", "Fecha Hasta",
    "Código Producto", "Producto", "Categoría",
    "Stock Inicial Anterior", "Stock Inicial Excel", "Alerta Continuidad",
    "Compras", "Promedio Compras Histórico", "Factor Desvío Compra",
    "Ventas usadas", "Ventas Delivery", "Otros Ingresos", "Otras Salidas", "Stock Final Excel",
    "Stock Esperado Sistema", "Venta Teórica", "Conteo Empleado", "Ajuste Admin",
    "Stock Final Físico", "Diferencia (VT - VR)", "Tipo Diferencia", "Severidad",
    "Costo Unitario", "Fuente Costo", "Impacto",
    "Cantidad Merma", "Cantidad Vencida", "Diferencia Anterior Compensada",
    "Posible Causa Principal", "Evidencia", "Nivel de Confianza",
    "Estado Auditoría", "Usuario Conteo", "Usuario Ajuste", "Fecha Ajuste",
    "Justificación Manual", "Observación", "Usuario Justificación", "Fecha Justificación",
]


def _excel_seguro(valor):
    """Neutraliza texto que Excel interpretaría como fórmula ejecutable."""
    if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def generar_excel_auditoria(periodo: InventarioPeriodo) -> io.BytesIO:
    """Devuelve el XLSX completo de un período listo para descargar."""
    resultados = (
        AuditoriaResultado.query
        .filter_by(periodo_id=periodo.id)
        .filter(AuditoriaResultado.estado_auditoria != "Archivado")
        .order_by(AuditoriaResultado.impacto.desc())
        .all()
    )

    # Defensa para auditorías antiguas: si una fila fue excluida después de la
    # última ejecución, tampoco debe aparecer en la descarga mientras se
    # re-ejecuta el período.
    excel = (
        ExcelImportado.query
        .filter_by(periodo_id=periodo.id, cliente_id=periodo.cliente_id)
        .filter(ExcelImportado.estado_validacion.in_(("ok", "pendiente_vinculacion")))
        .order_by(ExcelImportado.id.desc()).first()
    )
    if excel:
        filas_excel = ExcelDetalle.query.filter_by(excel_id=excel.id).all()
        nombres_activos = {
            (fila.producto_nombre_interno or fila.artdescrip or "").strip().casefold()
            for fila in filas_excel if not fila.excluido_auditoria
        }
        codigos_activos = {
            str(fila.articulo or "").strip().casefold()
            for fila in filas_excel if not fila.excluido_auditoria and fila.articulo
        }
        nombres_excluidos = {
            (fila.producto_nombre_interno or fila.artdescrip or "").strip().casefold()
            for fila in filas_excel if fila.excluido_auditoria
        } - nombres_activos
        codigos_excluidos = {
            str(fila.articulo or "").strip().casefold()
            for fila in filas_excel if fila.excluido_auditoria and fila.articulo
        } - codigos_activos
        resultados = [
            resultado for resultado in resultados
            if resultado.producto_nombre.strip().casefold() not in nombres_excluidos
            and str(resultado.articulo_codigo or "").strip().casefold() not in codigos_excluidos
        ]

    libro = openpyxl.Workbook()
    hoja = libro.active
    if hoja is None:
        hoja = libro.create_sheet()
    hoja.title = f"Auditoría Inv.{periodo.numero}"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")
    faltante_fill = PatternFill("solid", fgColor="FEE2E2")
    sobrante_fill = PatternFill("solid", fgColor="DCFCE7")
    critico_fill = PatternFill("solid", fgColor="FCA5A5")

    hoja.append(HEADERS_AUDITORIA)
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS_AUDITORIA))}1"
    for celda in hoja[1]:
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(horizontal="center")

    for resultado in resultados:
        justificacion = (
            Justificacion.query
            .filter_by(resultado_id=resultado.id)
            .order_by(Justificacion.id.desc())
            .first()
        )
        hoja.append([_excel_seguro(valor) for valor in [
            periodo.numero, periodo.fecha_desde, periodo.fecha_hasta,
            resultado.articulo_codigo, resultado.producto_nombre, resultado.categoria,
            resultado.stock_inicial_anterior, resultado.stock_inicial_excel,
            "SÍ" if resultado.alerta_continuidad else "No",
            resultado.compras, resultado.promedio_compras_historico,
            resultado.factor_desvio_compra,
            resultado.ventas, resultado.ventas_delivery, resultado.otros_ingresos,
            resultado.otras_salidas, resultado.stock_final_excel,
            resultado.stock_esperado, resultado.venta_teorica, resultado.conteo_empleado,
            resultado.ajuste_admin, resultado.conteo_final, resultado.diferencia,
            resultado.tipo_diferencia, resultado.severidad,
            resultado.costo_unitario, resultado.fuente_costo, resultado.impacto,
            resultado.cantidad_merma, resultado.cantidad_vencida,
            resultado.diferencia_anterior_compensada,
            resultado.causa_sugerida, resultado.evidencia,
            resultado.nivel_confianza, resultado.estado_auditoria,
            resultado.usuario_conteo, resultado.usuario_ajuste,
            resultado.fecha_ajuste,
            justificacion.causa if justificacion else "",
            justificacion.observacion if justificacion else "",
            justificacion.usuario if justificacion else "",
            str(justificacion.fecha.date()) if justificacion and justificacion.fecha else "",
        ]])
        fila = hoja.max_row
        fill = None
        if resultado.severidad == "Crítico":
            fill = critico_fill
        elif resultado.tipo_diferencia == "faltante":
            fill = faltante_fill
        elif resultado.tipo_diferencia == "sobrante":
            fill = sobrante_fill
        if fill:
            for celda in hoja[fila]:
                celda.fill = fill

    for columna in hoja.columns:
        max_len = max((len(str(celda.value)) for celda in columna if celda.value), default=10)
        if columna[0].column is None:
            continue
        letra = get_column_letter(columna[0].column)
        hoja.column_dimensions[letra].width = min(max_len + 4, 60)

    stream = io.BytesIO()
    libro.save(stream)
    libro.close()
    stream.seek(0)
    return stream
